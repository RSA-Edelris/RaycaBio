
import openpyxl, numpy as np, os
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

compound_order = ['BX912','EL2003A','EL2003A-A2U1','EL2003A-A4U1','EL5001A','EL5003A']

# Kinases as rows, compounds as columns, sorted by mean desc
tbl = pivot.T[compound_order].copy()
tbl['Mean'] = tbl.mean(axis=1)
tbl = tbl.sort_values('Mean', ascending=False)

def score_fill(val):
    if val is None or (isinstance(val, float) and np.isnan(val)):
        return PatternFill(fill_type=None)
    lo, hi = 6.3, 7.5
    t = max(0.0, min(1.0, (val - lo) / (hi - lo)))
    if t < 0.5:
        s = t * 2
        r = int(215 + (255-215)*(1-s)); g = int(48+(255-48)*s); b = int(39+(191-39)*s)
    else:
        s = (t-0.5)*2
        r = int(255+(26-255)*s); g = int(255+(152-255)*s); b = int(191+(80-191)*s)
    return PatternFill("solid", fgColor=f"{r:02X}{g:02X}{b:02X}")

hdr_fill  = PatternFill("solid", fgColor="1F3864")
mean_fill = PatternFill("solid", fgColor="2E4057")
hdr_font  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
body_font = Font(name="Calibri", size=10)
bold_font = Font(bold=True, name="Calibri", size=10)
ctr       = Alignment(horizontal="center", vertical="center")
thin      = Side(style="thin", color="CCCCCC")
brd       = Border(left=thin, right=thin, top=thin, bottom=thin)

wb    = openpyxl.Workbook()
ws_xl = wb.active
ws_xl.title = "Kinome Selectivity"

# Header
for ci, h in enumerate(["Kinase"] + compound_order + ["Mean pIC50"], 1):
    c = ws_xl.cell(1, ci, h); c.fill=hdr_fill; c.font=hdr_font; c.alignment=ctr; c.border=brd

# Data rows
for ri, (kinase, row) in enumerate(tbl.iterrows(), 2):
    c = ws_xl.cell(ri, 1, kinase); c.font=bold_font; c.alignment=ctr; c.border=brd
    for ci, cn in enumerate(compound_order, 2):
        v = row[cn]; dv = round(float(v),2) if not np.isnan(v) else None
        c = ws_xl.cell(ri, ci, dv)
        c.fill=score_fill(dv); c.font=body_font; c.alignment=ctr; c.border=brd; c.number_format="0.00"
    mv = round(float(row['Mean']),2)
    c = ws_xl.cell(ri, 8, mv)
    c.fill=mean_fill; c.font=Font(bold=True,color="FFFFFF",name="Calibri",size=10)
    c.alignment=ctr; c.border=brd; c.number_format="0.00"

ws_xl.column_dimensions['A'].width = 16
for ci in range(2, 9): ws_xl.column_dimensions[get_column_letter(ci)].width = 14
ws_xl.row_dimensions[1].height = 22
ws_xl.freeze_panes = "B2"

# Sheet 2: compounds as rows, kinases as columns
ws2 = wb.create_sheet("Raw (compounds × kinases)")
raw = pivot[sorted(pivot.columns)]
ws2.cell(1,1,"Compound").fill=hdr_fill; ws2.cell(1,1).font=hdr_font; ws2.cell(1,1).alignment=ctr
for ci, k in enumerate(raw.columns, 2):
    c=ws2.cell(1,ci,k); c.fill=hdr_fill; c.font=hdr_font; c.alignment=ctr
for ri, (cmpd, row) in enumerate(raw.iterrows(), 2):
    ws2.cell(ri,1,cmpd).font=bold_font
    for ci, v in enumerate(row, 2):
        dv = round(float(v),2) if not np.isnan(v) else None
        c=ws2.cell(ri,ci,dv); c.fill=score_fill(dv); c.font=body_font
        c.alignment=ctr; c.number_format="0.00"
ws2.column_dimensions['A'].width = 18
for ci in range(2, len(raw.columns)+2): ws2.column_dimensions[get_column_letter(ci)].width = 10
ws2.freeze_panes = "B2"

out = f"{WS}/kinome_selectivity_matrix.xlsx"
wb.save(out)
print(f"Saved: {out}  ({os.path.getsize(out):,} bytes)")
